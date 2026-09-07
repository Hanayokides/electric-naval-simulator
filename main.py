import os
import math
import pandas as pd
from datetime import datetime
import numpy as np
import matplotlib.pyplot as plt
from scipy.optimize import minimize_scalar
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from calculos_comuns import ModeloComumNaval
from recargas import RecargaRapida, RecargaMista, RecargaLenta
from diesel import SimuladorDiesel


class GerenciadorRelatorio:
    def __init__(self):
        self.comum = ModeloComumNaval(debug=True)
        self.rapida = RecargaRapida(self.comum)
        self.mista = RecargaMista(self.comum)
        self.lenta = RecargaLenta(self.comum)
        self.diesel = SimuladorDiesel(self.comum)

        self.velocidades_alvo = [8, 12, 16, 20]
        self.dods_m_lista = [12.5, 25.0, 37.5, 50.0, 62.5, 75.0, 90.0]
        self.distancias_rota_km = [5.0, 7.5, 10.0, 12.5, 15.0]
        self.potencias_mista_alvo = [2000, 2500, 3000, 3500, 4000, 4500, 5000]

        # Potências e DODs de comparação da Mista, específicos por terminal
        self.mista_1t_potencias_alvo = [3800, 3000, 2000, 1000, 500]
        self.mista_1t_dods_alvo = [90, 50, 43, 25]
        self.mista_2t_potencias_alvo = [1800, 1400, 1000, 600]
        self.mista_2t_dods_alvo = [90, 50, 37, 25]

        # Faixas de potência e DODs fixos para os gráficos TIR x Potência (4 por terminal, 8 no total)
        self.mista_1t_grafico_pot_range = (480, 3800)
        self.mista_2t_grafico_pot_range = (580, 1800)
        self.mista_1t_dods_grafico = [90, 50, 25]
        self.mista_2t_dods_grafico = [90, 50, 25]

    def rodar_otimizacao_rapida(self, L, tipo_terminal, v_knts, d_km):
        def objetivo(dod_tentativa):
            if dod_tentativa <= 1.0:
                return 1e9
            res = self.rapida.simular(L, dod_tentativa, tipo_terminal, v_knts, d_km=d_km)
            return 1e9 if (math.isnan(res["tir"]) or not res["convergido"]) else -res["tir"]
        return minimize_scalar(objetivo, bounds=(0.5, 20.0), method='bounded').x

    def rodar_otimizacao_lenta(self, L, v_knts, d_km):
        def objetivo(dod_tentativa):
            if dod_tentativa <= 1.0:
                return 1e9
            res = self.lenta.simular(L, dod_tentativa, "Lenta", v_knts, d_km=d_km)
            return 1e9 if (math.isnan(res["tir"]) or not res["convergido"]) else -res["tir"]
        return minimize_scalar(objetivo, bounds=(1.0, 3.0), method='bounded').x

    def rodar_otimizacao_mista(self, L, tipo_terminal, v_knts, d_km, pot_kw):
        def objetivo(dod_tentativa):
            if dod_tentativa <= 1.0:
                return 1e9
            res = self.mista.simular(L, dod_tentativa, tipo_terminal, v_knts, d_km=d_km, pot_infra_manual=pot_kw)
            return 1e9 if (math.isnan(res["tir"]) or not res["convergido"]) else -res["tir"]
        return minimize_scalar(objetivo, bounds=(1.0, 20.0), method='bounded').x

    def rodar_otimizacao_mista_potencia(self, L, tipo_terminal, v_knts, d_km, dod_pct, pot_min=200.0, pot_max=6000.0, n_pontos=120):
        # DOD fixo; varia a potência de infraestrutura para achar a que maximiza o TIR.
        # Busca em grade (a curva TIR x Potência pode ter trechos sem convergência).
        variavel_dod = 1.0 / (dod_pct / 100.0)
        melhor_pot, melhor_res, melhor_tir = None, None, -1e18
        for pot in np.linspace(pot_min, pot_max, n_pontos):
            res = self.mista.simular(L, variavel_dod, tipo_terminal, v_knts, d_km=d_km, pot_infra_manual=pot)
            if res["convergido"] and not math.isnan(res["tir"]) and res["tir"] > melhor_tir:
                melhor_tir = res["tir"]
                melhor_pot = pot
                melhor_res = res
        return melhor_pot, melhor_res

    def gerar_grafico_suave_dod(self, L, v_knts, d_km):
        dods_porcentagem = np.linspace(5, 100, 100)
        d_ef_axis, tirs_1t, tirs_2t, tirs_lenta = [], [], [], []

        for d in dods_porcentagem:
            v_dod = 1.0 / (max(0.05, d) / 100.0)
            res_1t = self.rapida.simular(L, v_dod, "1T", v_knts, d_km=d_km)
            res_2t = self.rapida.simular(L, v_dod, "2T", v_knts, d_km=d_km)
            res_len = self.lenta.simular(L, v_dod, "Lenta", v_knts, d_km=d_km)

            tirs_1t.append(res_1t["tir"] * 100 if (not math.isnan(res_1t["tir"]) and res_1t["convergido"]) else None)
            tirs_2t.append(res_2t["tir"] * 100 if (not math.isnan(res_2t["tir"]) and res_2t["convergido"]) else None)
            tirs_lenta.append(res_len["tir"] * 100 if (not math.isnan(res_len["tir"]) and res_len["convergido"]) else None)
            d_ef_axis.append(d / 100.0)

        plt.figure(figsize=(8.5, 5.5))
        plt.plot(d_ef_axis, tirs_1t, label='Fast Charging - 1 Terminal', color='#1f77b4', linewidth=2)
        plt.plot(d_ef_axis, tirs_2t, label='Fast Charging - 2 Terminals', color='#ff7f0e', linewidth=2)
        plt.plot(d_ef_axis, tirs_lenta, label='Slow Charging (Overnight)', color='#2ca02c', linewidth=2, linestyle='--')

        plt.title(f'Sensitivity: IRR vs Actual Effective DOD ({v_knts} knots - Route: {d_km}km)', fontsize=11, fontweight='bold', pad=12)
        plt.xlabel('Effective DOD (Fraction from 0.0 to 1.0)', fontsize=9)
        plt.ylabel('Internal Rate of Return IRR (%)', fontsize=9)
        plt.xlim(0.0, 1.0)

        valores_validos = [v for v in tirs_1t + tirs_2t + tirs_lenta if v is not None]
        if valores_validos:
            plt.ylim(min(valores_validos) - 2.0, max(valores_validos) + 2.0)

        plt.grid(True, linestyle='--', alpha=0.5)
        plt.legend(frameon=True, facecolor='white', loc='lower center')
        plt.tight_layout()
        path = f'grafico_efe_dod_{v_knts}knts_{d_km}km.png'
        plt.savefig(path, dpi=300)
        plt.close()
        return path

    def gerar_grafico_mista_potencia(self, L, v_knts, d_km):
        fisica_est = self.mista.calcular_demanda_energetica(L, 1.2, v_knts, d_km)
        energia_rota = fisica_est["energia_rota"]
        pot_max_teorica = (2 * energia_rota) / (5.0 / 60.0)

        pot_min = 200
        pot_max = math.ceil((pot_max_teorica * 1.1) / 100.0) * 100
        potencias = np.linspace(pot_min, pot_max, 100)
        tirs_m1t, tirs_m2t = [], []

        for p in potencias:
            dod_ot_m1t = self.rodar_otimizacao_mista(L, "1T", v_knts, d_km, p)
            res_m1t = self.mista.simular(L, dod_ot_m1t, "1T", v_knts, d_km=d_km, pot_infra_manual=p)
            tirs_m1t.append(res_m1t["tir"] * 100 if (res_m1t["convergido"] and not math.isnan(res_m1t["tir"])) else None)

            dod_ot_m2t = self.rodar_otimizacao_mista(L, "2T", v_knts, d_km, p)
            res_m2t = self.mista.simular(L, dod_ot_m2t, "2T", v_knts, d_km=d_km, pot_infra_manual=p)
            tirs_m2t.append(res_m2t["tir"] * 100 if (res_m2t["convergido"] and not math.isnan(res_m2t["tir"])) else None)

        plt.figure(figsize=(8.5, 5.5))
        v_m1t_p = [p for p, t in zip(potencias, tirs_m1t) if t is not None]
        v_m1t_t = [t for t in tirs_m1t if t is not None]
        if v_m1t_p:
            plt.plot(v_m1t_p, v_m1t_t, label='Mixed Charging 1T', color='#9467bd', linewidth=2.5)

        v_m2t_p = [p for p, t in zip(potencias, tirs_m2t) if t is not None]
        v_m2t_t = [t for t in tirs_m2t if t is not None]
        if v_m2t_p:
            plt.plot(v_m2t_p, v_m2t_t, label='Mixed Charging 2T', color='#8c564b', linewidth=2.5)

        plt.title(f'Mixed Charging Sensitivity: IRR vs Power ({v_knts} knots - Route: {d_km}km)', fontsize=11, fontweight='bold', pad=12)
        plt.xlabel('Installed Infrastructure Power (kW)', fontsize=9)
        plt.ylabel('Maximum Internal Rate of Return IRR (%)', fontsize=9)
        plt.xlim(pot_min, pot_max)
        plt.grid(True, linestyle='--', alpha=0.6)
        plt.legend(frameon=True, facecolor='white', loc='best')
        plt.tight_layout()
        path_viabilidade = f'grafico_mista_pot_{v_knts}knts_{d_km}km.png'
        plt.savefig(path_viabilidade, dpi=300)
        plt.close()
        return path_viabilidade

    def encontrar_dod_otimo_mista(self, L, tipo_terminal, v_knts, d_km, pot_referencia):
        # DOD que maximiza o TIR numa potência de referência. Busca em grade.
        melhor_dod, melhor_tir = 50.0, -1e18
        for dod_pct in np.linspace(10, 90, 100):
            res = self.mista.simular(L, 100.0 / dod_pct, tipo_terminal, v_knts, d_km=d_km, pot_infra_manual=pot_referencia)
            if res["convergido"] and not math.isnan(res["tir"]) and res["tir"] > melhor_tir:
                melhor_tir = res["tir"]
                melhor_dod = dod_pct
        return melhor_dod

    def gerar_grafico_mista_tir_potencia_dod_fixo(self, L, v_knts, d_km, tipo_terminal, dod_pct, pot_min, pot_max, label_dod=None):
        variavel_dod = 1.0 / (dod_pct / 100.0)
        potencias = np.linspace(pot_min, pot_max, 40)
        tirs = []
        for p in potencias:
            res = self.mista.simular(L, variavel_dod, tipo_terminal, v_knts, d_km=d_km, pot_infra_manual=p)
            tirs.append(res["tir"] * 100 if (not math.isnan(res["tir"]) and res["convergido"]) else None)

        cor = '#9467bd' if tipo_terminal == "1T" else '#8c564b'
        rotulo = label_dod if label_dod else f'DOD {dod_pct:.1f}%'

        plt.figure(figsize=(8.5, 5.5))
        pontos_x = [p for p, t in zip(potencias, tirs) if t is not None]
        pontos_y = [t for t in tirs if t is not None]
        if pontos_x:
            plt.plot(pontos_x, pontos_y, label=f'Mixed Charging {tipo_terminal} - {rotulo}', color=cor, linewidth=2.5)

        plt.title(f'Mixed Charging Sensitivity: IRR vs Power ({tipo_terminal}, {rotulo} | {v_knts} knots - Route: {d_km}km)', fontsize=11, fontweight='bold', pad=12)
        plt.xlabel('Installed Infrastructure Power (kW)', fontsize=9)
        plt.ylabel('Internal Rate of Return - IRR (%)', fontsize=9)
        plt.xlim(pot_min, pot_max)

        if pontos_y:
            plt.ylim(min(pontos_y) - 2.0, max(pontos_y) + 2.0)

        plt.grid(True, linestyle='--', alpha=0.5)
        plt.legend(frameon=True, facecolor='white', loc='best')
        plt.tight_layout()
        dod_tag = f"{dod_pct:.0f}".replace('.', '_')
        path = f'grafico_mista_potfixa_dod{dod_tag}_{tipo_terminal}_{v_knts}knts_{d_km}km.png'
        plt.savefig(path, dpi=300)
        plt.close()
        return path

    def exportar_para_excel(self, L, resultados_excel, timestamp_str):
        rows_data = []

        def fmt_vpl(r):
            return float(r['vpl']) if (r["convergido"] and not math.isnan(r['vpl'])) else "Não Convergiu"

        def fmt_tir(r):
            return float(r['tir']) * 100 if (r["convergido"] and not math.isnan(r['tir'])) else "NaN"

        def fmt_num(r, k):
            if r.get(k) == "-":
                return "-"
            return float(r[k]) if r["convergido"] else "N/A"

        colunas_principais = [
            "Tamanho da Rota (km)", "Configuração do Cenário",
            "Quantidade de Terminais / Modo", "Velocidade (knts)",
            "Tipo de Solução", "Variável DOD", "DOD Efetivo Diário (Real)", "DOD Alvo",
            "Taxa Interna de Retorno (TIR)", "Valor Presente Liquido (VPL) USD",
            "Calado (T) (m)", "Massa Total Barcaça (kg)",
            "Deslocamento Empuxo (kg)", "Energia Banco Bateria (kWh)",
            "Custo Banco Bateria USD", "Custos de Infraestrutura USD",
            "Custo Anual Troca Bateria USD", "Pot. Infraestrutura Recarga (kW)",
            "C-rate Carga", "C-rate Descarga", "Número de Trocas/Ano",
            "Número de Ciclos EOL", "Potência Nominal Eixo (kW)"
        ]

        chaves_ignoradas = {"tir", "vpl", "calado", "massa_total_barca", "deslocamento", "pot_infra",
                             "b_bateria_rec", "b_baterias", "trocas", "custo_baterias_anual_troca",
                             "infraestrutura_recarga", "crate_carga", "crate_descarga", "ciclos_eol",
                             "dod_efetivo", "pot_nominal", "convergido", "log_trace"}

        chaves_extras = set()
        for item in resultados_excel:
            if "res" in item and isinstance(item["res"], dict):
                for k in item["res"].keys():
                    if k not in chaves_ignoradas:
                        chaves_extras.add(k)

        chaves_extras = sorted(list(chaves_extras))
        colunas = colunas_principais + [" "] + chaves_extras

        for d_km in self.distancias_rota_km:
            rows_data.append({k: "" for k in colunas})
            rows_data.append({"Tamanho da Rota (km)": f"--- SEÇÃO ROTA: {d_km:.1f} km ---"})
            dados_filtrados = [x for x in resultados_excel if x["distancia"] == d_km]

            def sort_key(item):
                v = item.get("v", 0)
                t = item.get("t", "")
                sol_tipo = str(item.get("sol_tipo", ""))

                if t == "1T":
                    tipo_order = 1
                    pot = 0
                elif t == "2T":
                    tipo_order = 2
                    pot = 0
                elif t == "Lenta":
                    tipo_order = 3
                    pot = 0
                elif str(t).startswith("M1T"):
                    tipo_order = 4
                    try:
                        pot = float(t.split("_")[1])
                    except Exception:
                        pot = 0
                elif str(t).startswith("M2T"):
                    tipo_order = 5
                    try:
                        pot = float(t.split("_")[1])
                    except Exception:
                        pot = 0
                elif t == "Diesel":
                    tipo_order = 6
                    pot = 0
                else:
                    tipo_order = 99
                    pot = 0

                if "DOD" in sol_tipo:
                    try:
                        dod_val = float(sol_tipo.replace("DOD", "").replace("%", "").strip())
                    except Exception:
                        dod_val = 0.0
                elif sol_tipo == "Ótimo" or sol_tipo == "Otimo":
                    dod_val = 999.0
                else:
                    dod_val = 0.0

                return (v, tipo_order, dod_val, pot)

            dados_filtrados.sort(key=sort_key)
            for item in dados_filtrados:
                nome_modo = f"{item['t']} Rápida" if item['t'] in ["1T", "2T"] else item['t']

                row_dict = {
                    "Tamanho da Rota (km)": d_km,
                    "Configuração do Cenário": f"{item['v']} knts ({nome_modo}) - {item['sol_tipo']}",
                    "Quantidade de Terminais / Modo": item['t'],
                    "Velocidade (knts)": item['v'],
                    "Tipo de Solução": "DOD Ótimo" if item['sol_tipo'] in ["Ótimo", "Otimo"] else item['sol_tipo'],
                    "Variável DOD": round(item['input'], 4) if isinstance(item['input'], (int, float)) else item['input'],
                    "DOD Efetivo Diário (Real)": fmt_num(item['res'], 'dod_efetivo'),
                    "DOD Alvo": fmt_num(item['res'], 'dod_alvo'),
                    "Taxa Interna de Retorno (TIR)": fmt_tir(item['res']),
                    "Valor Presente Liquido (VPL) USD": fmt_vpl(item['res']),
                    "Calado (T) (m)": fmt_num(item['res'], 'calado'),
                    "Massa Total Barcaça (kg)": fmt_num(item['res'], 'massa_total_barca'),
                    "Deslocamento Empuxo (kg)": fmt_num(item['res'], 'deslocamento'),
                    "Energia Banco Bateria (kWh)": fmt_num(item['res'], 'b_bateria_rec'),
                    "Custo Banco Bateria USD": fmt_num(item['res'], 'b_baterias'),
                    "Custo Anual Troca Bateria USD": fmt_num(item['res'], 'custo_baterias_anual_troca'),
                    "Custos de Infraestrutura USD": fmt_num(item['res'], 'infraestrutura_recarga'),
                    "Pot. Infraestrutura Recarga (kW)": fmt_num(item['res'], 'pot_infra'),
                    "C-rate Carga": fmt_num(item['res'], 'crate_carga'),
                    "C-rate Descarga": fmt_num(item['res'], 'crate_descarga'),
                    "Número de Trocas/Ano": fmt_num(item['res'], 'trocas'),
                    "Número de Ciclos EOL": fmt_num(item['res'], 'ciclos_eol'),
                    "Potência Nominal Eixo (kW)": fmt_num(item['res'], 'pot_nominal'),
                    " ": ""
                }

                for k in chaves_extras:
                    val = item['res'].get(k, "N/A")
                    if isinstance(val, (int, float)):
                        row_dict[k] = round(val, 6)
                    else:
                        row_dict[k] = val

                rows_data.append(row_dict)

        df = pd.DataFrame(rows_data, columns=colunas)
        filename = f"planilha_dimensionamento_completa_{L:.1f}m_{timestamp_str}.xlsx"
        df.to_excel(filename, index=False)
        wb = openpyxl.load_workbook(filename)
        ws = wb.active
        ws.views.sheetView[0].showGridLines = True
        header_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid")
        extra_header_fill = PatternFill(start_color="138D75", end_color="138D75", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        section_fill = PatternFill(start_color="2E86C1", end_color="2E86C1", fill_type="solid")
        section_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        zebra_fill = PatternFill(start_color="F2F9FF", end_color="F2F9FF", fill_type="solid")
        thin_border = Border(left=Side(style='thin', color='D3D3D3'), right=Side(style='thin', color='D3D3D3'),
                              top=Side(style='thin', color='D3D3D3'), bottom=Side(style='thin', color='D3D3D3'))

        for cell in ws[1]:
            if cell.value == " ":
                cell.fill = PatternFill(fill_type=None)
            elif str(cell.value) in chaves_extras:
                cell.fill = extra_header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28

        col_vazia_idx = colunas.index(" ") + 1
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            first_cell_val = str(row[0].value or "")
            if first_cell_val.startswith("--- SEÇÃO ROTA"):
                ws.row_dimensions[row_idx].height = 22
                for cell in row:
                    if cell.column == col_vazia_idx:
                        continue
                    cell.fill = section_fill
                    cell.font = section_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            elif first_cell_val == "":
                ws.row_dimensions[row_idx].height = 10
            else:
                ws.row_dimensions[row_idx].height = 20
                for cell in row:
                    if cell.column == col_vazia_idx:
                        continue
                    if row_idx % 2 == 0:
                        cell.fill = zebra_fill
                    cell.border = thin_border
                    if isinstance(cell.value, (int, float)):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="center", vertical="center")

        for col in ws.columns:
            if str(col[0].value) == " ":
                ws.column_dimensions[get_column_letter(col[0].column)].width = 4
                continue
            max_len = max([len(str(cell.value or "")) for cell in col if not str(cell.value or "").startswith("--- SEÇÃO")])
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 4, 12)

        wb.save(filename)
        print(f" [SUCESSO EXCEL] Planilha consolidada gerada em: {filename}")

    def adicionar_legenda(self, doc, texto, is_table=False, contador=None):
        prefixo = "Table" if is_table else "Figure"
        contador[0] += 1
        p = doc.add_paragraph()
        r = p.add_run(f"{prefixo} {contador[0]}. {texto}")
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.name = 'Calibri'
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)

    def adicionar_tabela_simples(self, doc, headers, rows):
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Light Grid Accent 1'
        hdr = table.rows[0].cells
        for i, htext in enumerate(headers):
            hdr[i].text = str(htext)
            for p in hdr[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.bold = True
                    r.font.size = Pt(9.5)
        for row in rows:
            cells = table.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = str(val)
                for p in cells[i].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return table

    def gerar_secao_resumo_executivo(self, doc, resultados_otimos_apenas, L_usuario):
        fig_count = [0]
        tab_count = [0]

        lookup = {}
        for item in resultados_otimos_apenas:
            lookup[(item['distancia'], item['v'], item['t'])] = item['res']

        rotas = self.distancias_rota_km
        velocidades = self.velocidades_alvo
        d_ref = 5.0 if 5.0 in rotas else rotas[0]
        v_ref = 12 if 12 in velocidades else velocidades[0]

        todas_config_eletricas = (["1T", "2T", "Lenta"]
                                   + [f"M1T_{p}" for p in self.mista_1t_potencias_alvo]
                                   + [f"M2T_{p}" for p in self.mista_2t_potencias_alvo])

        def tir_valido(res):
            if res is None or not res.get("convergido"):
                return None
            tir = res.get("tir")
            if tir is None or (isinstance(tir, float) and math.isnan(tir)):
                return None
            return tir * 100.0

        doc.add_paragraph("Summary Results").runs[0].font.bold = True

        # Figura 1 - IRR vs DOD no cenario de referencia (1T, 2T, Lenta)
        if d_ref in rotas and v_ref in velocidades:
            p_base = self.gerar_grafico_suave_dod(L_usuario, v_ref, d_ref)
            doc.add_picture(p_base, width=Inches(6.2))
            self.adicionar_legenda(
                doc,
                f"Internal Rate of Return as a function of depth of discharge (DOD) for fast charging "
                f"(1 and 2 terminals) and slow (overnight) charging, at the reference scenario "
                f"({v_ref} knots, {d_ref:.1f} km route).",
                contador=fig_count
            )
            os.remove(p_base)

        # Figura 2 - barras comparando todas as configuracoes no cenario de referencia
        labels, tirs, colors = [], [], []
        label_map = {"1T": "Fast\n1 Terminal", "2T": "Fast\n2 Terminals", "Lenta": "Slow\n(overnight)"}
        for t in ["1T", "2T", "Lenta"]:
            labels.append(label_map[t])
            tirs.append(tir_valido(lookup.get((d_ref, v_ref, t))))
            colors.append("#1f77b4")
        for p in self.mista_1t_potencias_alvo:
            labels.append(f"Mixed 1T\n{p} kW")
            tirs.append(tir_valido(lookup.get((d_ref, v_ref, f"M1T_{p}"))))
            colors.append("#9467bd")
        for p in self.mista_2t_potencias_alvo:
            labels.append(f"Mixed 2T\n{p} kW")
            tirs.append(tir_valido(lookup.get((d_ref, v_ref, f"M2T_{p}"))))
            colors.append("#8c564b")
        diesel_tir = tir_valido(lookup.get((d_ref, v_ref, "Diesel")))

        plt.figure(figsize=(10, 5))
        xs = range(len(labels))
        alturas = [t if t is not None else 0 for t in tirs]
        plt.bar(xs, alturas, color=colors)
        if diesel_tir is not None:
            plt.axhline(diesel_tir, color='black', linestyle='--', linewidth=1.3, label=f'Diesel baseline ({diesel_tir:.1f}%)')
            plt.legend(loc='upper right')
        plt.xticks(list(xs), labels, fontsize=8)
        plt.ylabel('Internal Rate of Return (%)')
        plt.title(f'IRR by charging strategy - reference scenario ({v_ref} knots, {d_ref:.1f} km route)')
        plt.grid(axis='y', linestyle='--', alpha=0.4)
        for i, v in enumerate(tirs):
            if v is not None:
                plt.text(i, v + (1.2 if v >= 0 else -2.2), f'{v:.1f}', ha='center', fontsize=8)
            else:
                plt.text(i, 0.5, 'n/c', ha='center', fontsize=8, color='gray', rotation=90)
        plt.tight_layout()
        path_bar = f'fig_resumo_barras_{v_ref}kn_{d_ref}km.png'
        plt.savefig(path_bar, dpi=300, bbox_inches='tight')
        plt.close()
        doc.add_picture(path_bar, width=Inches(6.3))
        self.adicionar_legenda(
            doc,
            f"Internal Rate of Return by charging strategy for the reference scenario "
            f"({v_ref} knots, {d_ref:.1f} km route). \"n/c\" denotes configurations for which no "
            f"converged, financially defined solution was obtained.",
            contador=fig_count
        )
        os.remove(path_bar)

        # Figura 3 - mapas de calor Diesel vs melhor eletrica em toda a grade
        diesel_grid = np.full((len(rotas), len(velocidades)), np.nan)
        eletrica_grid = np.full((len(rotas), len(velocidades)), np.nan)
        for i, r in enumerate(rotas):
            for j, v in enumerate(velocidades):
                dtir = tir_valido(lookup.get((r, v, "Diesel")))
                if dtir is not None:
                    diesel_grid[i, j] = dtir
                melhor = None
                for t in todas_config_eletricas:
                    etir = tir_valido(lookup.get((r, v, t)))
                    if etir is not None and (melhor is None or etir > melhor):
                        melhor = etir
                if melhor is not None:
                    eletrica_grid[i, j] = melhor

        fig, axes = plt.subplots(1, 2, figsize=(10, 4.2), constrained_layout=True)
        validos = np.concatenate([diesel_grid[~np.isnan(diesel_grid)], eletrica_grid[~np.isnan(eletrica_grid)]])
        vmin, vmax = (validos.min(), validos.max()) if len(validos) else (0, 1)
        im = None
        for ax, grid, title in zip(axes, [diesel_grid, eletrica_grid], ['Diesel baseline', 'Best electric configuration']):
            im = ax.imshow(grid, cmap='RdYlGn', vmin=vmin, vmax=vmax, aspect='auto')
            ax.set_xticks(range(len(velocidades))); ax.set_xticklabels(velocidades)
            ax.set_yticks(range(len(rotas))); ax.set_yticklabels(rotas)
            ax.set_xlabel('Speed (knots)'); ax.set_ylabel('Route length (km)')
            ax.set_title(title)
            for i in range(len(rotas)):
                for j in range(len(velocidades)):
                    val = grid[i, j]
                    if not np.isnan(val):
                        ax.text(j, i, f'{val:.1f}', ha='center', va='center', fontsize=8,
                                 color='black' if abs(val) < 25 else 'white')
                    else:
                        ax.text(j, i, 'n/c', ha='center', va='center', fontsize=8, color='gray')
        fig.colorbar(im, ax=axes, shrink=0.85, pad=0.02, label='Internal Rate of Return (%)')
        path_heat = 'fig_resumo_heatmaps.png'
        plt.savefig(path_heat, dpi=300, bbox_inches='tight')
        plt.close()
        doc.add_picture(path_heat, width=Inches(6.3))
        self.adicionar_legenda(
            doc,
            "Internal Rate of Return (%) across the tested route-length/speed envelope for the diesel "
            "baseline (left) and the best-performing electric configuration at each grid point (right). "
            "\"n/c\" denotes combinations for which no economically or numerically valid solution was obtained.",
            contador=fig_count
        )
        os.remove(path_heat)

        # Figura 4 - sensibilidade do TIR ao tamanho da rota, na velocidade de referencia
        plt.figure(figsize=(7.5, 5))
        series = [("Diesel", "Diesel", "black", "o", "-"), ("1T", "1T", "#1f77b4", "s", "-"),
                  ("2T", "2T", "#ff7f0e", "^", "-"), ("Lenta", "Slow (overnight)", "#2ca02c", "d", "--")]
        for t, label, color, marker, ls in series:
            xs_r, ys_r = [], []
            for r in rotas:
                v = tir_valido(lookup.get((r, v_ref, t)))
                if v is not None:
                    xs_r.append(r); ys_r.append(v)
            if xs_r:
                plt.plot(xs_r, ys_r, marker=marker, linestyle=ls, color=color, label=label, linewidth=1.8, markersize=6)
        plt.axhline(0, color='gray', linewidth=0.8)
        plt.xlabel('Route length (km)'); plt.ylabel('Internal Rate of Return (%)')
        plt.title(f'IRR sensitivity to route length at {v_ref} knots')
        plt.grid(True, linestyle='--', alpha=0.4)
        plt.legend()
        plt.tight_layout()
        path_sens = 'fig_resumo_sensibilidade_rota.png'
        plt.savefig(path_sens, dpi=300, bbox_inches='tight')
        plt.close()
        doc.add_picture(path_sens, width=Inches(5.8))
        self.adicionar_legenda(
            doc,
            f"Internal Rate of Return sensitivity to route length at a fixed cruising speed of "
            f"{v_ref} knots, for the diesel baseline and the three primary electric configurations.",
            contador=fig_count
        )
        os.remove(path_sens)

        # Figura 5 - DOD otimo em funcao da velocidade, por rota, para 1T e 2T
        fig, axes = plt.subplots(1, 2, figsize=(10, 4.3), sharey=True, constrained_layout=True)
        cores_rota = plt.cm.viridis_r([i / max(1, len(rotas) - 1) for i in range(len(rotas))])
        for ax, t, title in zip(axes, ["1T", "2T"], ["Fast charging - 1 terminal", "Fast charging - 2 terminals"]):
            for r, cor in zip(rotas, cores_rota):
                xs_v, ys_v = [], []
                for v in velocidades:
                    res = lookup.get((r, v, t))
                    if res is None or not res.get("convergido"):
                        continue
                    dod = res.get("dod_efetivo")
                    tir = res.get("tir")
                    if dod in (None, "-") or (isinstance(tir, float) and math.isnan(tir)) or dod <= 0.06:
                        continue
                    xs_v.append(v); ys_v.append(dod * 100.0)
                if xs_v:
                    ax.plot(xs_v, ys_v, marker='o', color=cor, label=f'{r:.1f} km')
            ax.axhspan(35, 40, color='gray', alpha=0.15)
            ax.set_xlabel('Speed (knots)')
            ax.set_title(title)
            ax.grid(True, linestyle='--', alpha=0.4)
        axes[0].set_ylabel('Optimal design DOD (%)')
        axes[1].legend(title='Route length', loc='upper right', fontsize=8)
        path_dod = 'fig_resumo_dod_otimo.png'
        plt.savefig(path_dod, dpi=300, bbox_inches='tight')
        plt.close()
        doc.add_picture(path_dod, width=Inches(6.3))
        self.adicionar_legenda(
            doc,
            "Optimal design DOD as a function of speed, for each route length, for single-terminal "
            "(left) and two-terminal (right) fast charging. The shaded band marks the 35-40% range. "
            "Non-convergent or non-viable points are excluded.",
            contador=fig_count
        )
        os.remove(path_dod)

        # Tabela 1 - resumo do cenario de referencia
        linhas_tab1 = []
        nomes = [("Diesel baseline", "Diesel"), ("Fast charging, 1 terminal", "1T"),
                 ("Fast charging, 2 terminals", "2T"), ("Slow (overnight) charging", "Lenta")]
        for nome, t in nomes:
            res = lookup.get((d_ref, v_ref, t))
            if res is None:
                continue
            tir = tir_valido(res)
            dod = res.get("dod_efetivo")
            dod_str = f"{dod*100:.1f}%" if isinstance(dod, (int, float)) else "-"
            banco = res.get("b_bateria_rec")
            banco_str = f"{banco:,.0f}" if isinstance(banco, (int, float)) else "-"
            pot = res.get("pot_infra")
            pot_str = f"{pot:,.0f}" if isinstance(pot, (int, float)) else "-"
            vpl = res.get("vpl")
            vpl_str = f"{vpl:,.0f}" if isinstance(vpl, (int, float)) and not math.isnan(vpl) else "N/A"
            tir_str = f"{tir:.1f}" if tir is not None else "N/A"
            linhas_tab1.append([nome, dod_str, banco_str, pot_str, tir_str, vpl_str])
        self.adicionar_tabela_simples(
            doc,
            ["Configuration", "Design DOD", "Battery capacity (kWh)", "Charger power (kW)", "IRR (%)", "NPV (USD)"],
            linhas_tab1
        )
        self.adicionar_legenda(
            doc,
            f"Optimal-DOD results for the reference scenario ({v_ref} knots, {d_ref:.1f} km route).",
            is_table=True, contador=tab_count
        )

        # Tabela 2 - velocidade maxima viavel (TIR >= TMA) por tamanho de rota
        tma = 7.0
        linhas_tab2 = []
        for r in rotas:
            v_max_diesel, v_max_eletrica = None, None
            for v in velocidades:
                dtir = tir_valido(lookup.get((r, v, "Diesel")))
                if dtir is not None and dtir >= tma:
                    v_max_diesel = v
                melhor = None
                for t in todas_config_eletricas:
                    etir = tir_valido(lookup.get((r, v, t)))
                    if etir is not None and (melhor is None or etir > melhor):
                        melhor = etir
                if melhor is not None and melhor >= tma:
                    v_max_eletrica = v
            linhas_tab2.append([
                f"{r:.1f}",
                str(v_max_diesel) if v_max_diesel is not None else "not viable",
                str(v_max_eletrica) if v_max_eletrica is not None else "not viable"
            ])
        self.adicionar_tabela_simples(
            doc,
            ["Route length (km)", "Max. viable speed - Diesel (kn)", "Max. viable speed - Electric, best config. (kn)"],
            linhas_tab2
        )
        self.adicionar_legenda(
            doc,
            f"Highest tested service speed at which each propulsion type remains a financially viable "
            f"investment (IRR >= {tma:.0f}%), by route length.",
            is_table=True, contador=tab_count
        )
        doc.add_paragraph("\n")

    def adicionar_borda_tabela(self, cell):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            r'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            r'<w:top w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
            r'<w:left w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
            r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
            r'<w:right w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
            r'</w:tcBorders>'
        )
        tcPr.append(borders)

    def executar(self):
        print("=" * 85)
        print("SISTEMA NAVAL PARAMÉTRICO: RECARGA RÁPIDA (1T/2T), MISTA (1T/2T) E LENTA")
        print("=" * 85)

        L_usuario = 53.0
        modo = None
        while modo is None:
            escolha = input("Digite M para múltiplos resultados, O para resultados ótimos e ESC para sair (fechar o código): ").strip().upper()
            if escolha in ['M', 'O']:
                modo = escolha
            elif escolha == 'ESC':
                print("Encerrando o programa.")
                return
            else:
                print("Digite uma opção possível.")

        timestamp_agora = datetime.now().strftime("%d.%m.%Y_%H%M")
        resultados_mestre, resultados_otimos_apenas = [], []

        for d_km in self.distancias_rota_km:
            print(f" [PROCESSO] Simulando Rota de {d_km} km...")
            for v in self.velocidades_alvo:
                v_ot_1t = self.rodar_otimizacao_rapida(L_usuario, "1T", v, d_km)
                v_ot_2t = self.rodar_otimizacao_rapida(L_usuario, "2T", v, d_km)
                v_ot_len = self.rodar_otimizacao_lenta(L_usuario, v, d_km)

                r_1t = self.rapida.simular(L_usuario, v_ot_1t, "1T", v, d_km=d_km)
                r_2t = self.rapida.simular(L_usuario, v_ot_2t, "2T", v, d_km=d_km)
                r_len = self.lenta.simular(L_usuario, v_ot_len, "Lenta", v, d_km=d_km)
                r_diesel = self.diesel.simular(L_usuario, v, d_km=d_km)

                for t, ot, r in [("1T", v_ot_1t, r_1t), ("2T", v_ot_2t, r_2t), ("Lenta", v_ot_len, r_len), ("Diesel", "-", r_diesel)]:
                    tipo_sol = "Baseline Combustão" if t == "Diesel" else "Ótimo"
                    resultados_mestre.append({"distancia": d_km, "v": v, "t": t, "sol_tipo": tipo_sol, "input": ot, "res": r})
                    resultados_otimos_apenas.append({"distancia": d_km, "v": v, "t": t, "res": r})

                for p_mista in self.mista_1t_potencias_alvo:
                    v_ot_m1t = self.rodar_otimizacao_mista(L_usuario, "1T", v, d_km, p_mista)
                    r_m1t = self.mista.simular(L_usuario, v_ot_m1t, "1T", v, d_km=d_km, pot_infra_manual=p_mista)
                    resultados_mestre.append({"distancia": d_km, "v": v, "t": f"M1T_{p_mista}", "sol_tipo": "Ótimo", "input": v_ot_m1t, "res": r_m1t})
                    resultados_otimos_apenas.append({"distancia": d_km, "v": v, "t": f"M1T_{p_mista}", "res": r_m1t})

                for p_mista in self.mista_2t_potencias_alvo:
                    v_ot_m2t = self.rodar_otimizacao_mista(L_usuario, "2T", v, d_km, p_mista)
                    r_m2t = self.mista.simular(L_usuario, v_ot_m2t, "2T", v, d_km=d_km, pot_infra_manual=p_mista)
                    resultados_mestre.append({"distancia": d_km, "v": v, "t": f"M2T_{p_mista}", "sol_tipo": "Ótimo", "input": v_ot_m2t, "res": r_m2t})
                    resultados_otimos_apenas.append({"distancia": d_km, "v": v, "t": f"M2T_{p_mista}", "res": r_m2t})

                if modo == 'M':
                    for d in self.dods_m_lista:
                        v_conv = 1.0 / (d / 100.0)
                        for t, r_func in [("1T", self.rapida), ("2T", self.rapida), ("Lenta", self.lenta)]:
                            r_m = r_func.simular(L_usuario, v_conv, t, v, d_km=d_km)
                            resultados_mestre.append({"distancia": d_km, "v": v, "t": t, "sol_tipo": f"DOD {d}%", "input": v_conv, "res": r_m})

                    for d in self.mista_1t_dods_alvo:
                        v_conv = 1.0 / (d / 100.0)
                        for p_mista in self.mista_1t_potencias_alvo:
                            r_m_m1t = self.mista.simular(L_usuario, v_conv, "1T", v, d_km=d_km, pot_infra_manual=p_mista)
                            resultados_mestre.append({"distancia": d_km, "v": v, "t": f"M1T_{p_mista}", "sol_tipo": f"DOD {d}%", "input": v_conv, "res": r_m_m1t})

                    for d in self.mista_2t_dods_alvo:
                        v_conv = 1.0 / (d / 100.0)
                        for p_mista in self.mista_2t_potencias_alvo:
                            r_m_m2t = self.mista.simular(L_usuario, v_conv, "2T", v, d_km=d_km, pot_infra_manual=p_mista)
                            resultados_mestre.append({"distancia": d_km, "v": v, "t": f"M2T_{p_mista}", "sol_tipo": f"DOD {d}%", "input": v_conv, "res": r_m_m2t})

                    for d in self.mista_1t_dods_alvo:
                        pot_ot, res_ot = self.rodar_otimizacao_mista_potencia(L_usuario, "1T", v, d_km, d)
                        if res_ot is not None:
                            resultados_mestre.append({"distancia": d_km, "v": v, "t": "M1T_PotOtima", "sol_tipo": f"DOD {d}%", "input": 1.0 / (d / 100.0), "res": res_ot})

                    for d in self.mista_2t_dods_alvo:
                        pot_ot, res_ot = self.rodar_otimizacao_mista_potencia(L_usuario, "2T", v, d_km, d)
                        if res_ot is not None:
                            resultados_mestre.append({"distancia": d_km, "v": v, "t": "M2T_PotOtima", "sol_tipo": f"DOD {d}%", "input": 1.0 / (d / 100.0), "res": res_ot})

        self.exportar_para_excel(L_usuario, resultados_mestre, timestamp_agora)

        doc = Document()
        doc.add_paragraph("HYDRODYNAMIC-FINANCIAL SENSITIVITY REPORT").runs[0].font.bold = True

        self.gerar_secao_resumo_executivo(doc, resultados_otimos_apenas, L_usuario)

        for d_km in self.distancias_rota_km:
            doc.add_paragraph(f"Operational Scenario Route Length: {d_km:.1f} km").runs[0].font.bold = True
            for v in self.velocidades_alvo:
                p_grafico1 = self.gerar_grafico_suave_dod(L_usuario, v, d_km)
                doc.add_picture(p_grafico1, width=Inches(6.2))
                os.remove(p_grafico1)

                pot_min_1t, pot_max_1t = self.mista_1t_grafico_pot_range
                dod_otimo_1t = self.encontrar_dod_otimo_mista(L_usuario, "1T", v, d_km, pot_referencia=(pot_min_1t + pot_max_1t) / 2)
                for dod_pct in self.mista_1t_dods_grafico:
                    p_g_mista = self.gerar_grafico_mista_tir_potencia_dod_fixo(L_usuario, v, d_km, "1T", dod_pct, pot_min_1t, pot_max_1t)
                    doc.add_picture(p_g_mista, width=Inches(6.2))
                    os.remove(p_g_mista)
                p_g_mista = self.gerar_grafico_mista_tir_potencia_dod_fixo(L_usuario, v, d_km, "1T", dod_otimo_1t, pot_min_1t, pot_max_1t, label_dod=f"DOD Ótimo ({dod_otimo_1t:.1f}%)")
                doc.add_picture(p_g_mista, width=Inches(6.2))
                os.remove(p_g_mista)

                pot_min_2t, pot_max_2t = self.mista_2t_grafico_pot_range
                dod_otimo_2t = self.encontrar_dod_otimo_mista(L_usuario, "2T", v, d_km, pot_referencia=(pot_min_2t + pot_max_2t) / 2)
                for dod_pct in self.mista_2t_dods_grafico:
                    p_g_mista = self.gerar_grafico_mista_tir_potencia_dod_fixo(L_usuario, v, d_km, "2T", dod_pct, pot_min_2t, pot_max_2t)
                    doc.add_picture(p_g_mista, width=Inches(6.2))
                    os.remove(p_g_mista)
                p_g_mista = self.gerar_grafico_mista_tir_potencia_dod_fixo(L_usuario, v, d_km, "2T", dod_otimo_2t, pot_min_2t, pot_max_2t, label_dod=f"DOD Ótimo ({dod_otimo_2t:.1f}%)")
                doc.add_picture(p_g_mista, width=Inches(6.2))
                os.remove(p_g_mista)
            doc.add_paragraph("\n")

        doc.add_paragraph("DATA MATRIX OPTIMAL DOD BY CONFIGURATION").runs[0].font.bold = True
        df_tabela = pd.DataFrame(resultados_otimos_apenas)

        def set_cell_bg(cell, hex_color):
            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{hex_color}"/>')
            tcPr.append(shd)

        def criar_tabela_config(id_tabela, titulo_sub):
            doc.add_paragraph(f"Configuration: {titulo_sub}").runs[0].font.bold = True
            headers = [f"Route/Speed"] + [f"{v} knots" for v in self.velocidades_alvo]
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            for i, text in enumerate(headers):
                hdr_cells[i].text = text
                p = hdr_cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.runs[0]
                run.font.bold, run.font.size, run.font.color.rgb = True, Pt(10), RGBColor(255, 255, 255)
                set_cell_bg(hdr_cells[i], "1B4F72")
                self.adicionar_borda_tabela(hdr_cells[i])

            for d_km in self.distancias_rota_km:
                row_cells = table.add_row().cells
                d_str = f"{d_km:.1f}".replace('.', ',')
                if d_str.endswith(',0'):
                    d_str = d_str[:-2]
                row_cells[0].text = d_str
                p0 = row_cells[0].paragraphs[0]
                p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run0 = p0.runs[0]
                run0.font.bold, run0.font.size = True, Pt(9.5)
                set_cell_bg(row_cells[0], "D6EAF8")
                self.adicionar_borda_tabela(row_cells[0])

                for col_idx, v in enumerate(self.velocidades_alvo, start=1):
                    res_lista = df_tabela[(df_tabela['distancia'] == d_km) & (df_tabela['v'] == v) & (df_tabela['t'] == id_tabela)]['res'].values
                    val_str = "N/A"
                    if len(res_lista) > 0:
                        r = res_lista[0]
                        if not math.isnan(r['tir']) and r['convergido']:
                            val_dod = r.get('dod_efetivo', 0)
                            if val_dod == "-":
                                val_str = "-"
                            else:
                                val_dod = val_dod * 100.0
                                val_str = f"{val_dod:.1f}".replace('.', ',')
                                if val_str.endswith(',0'):
                                    val_str = val_str[:-2]

                    row_cells[col_idx].text = val_str
                    p_cell = row_cells[col_idx].paragraphs[0]
                    p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cell.runs[0].font.size = Pt(9.5)
                    self.adicionar_borda_tabela(row_cells[col_idx])
            doc.add_paragraph("\n")

        criar_tabela_config("1T", "Fast Charging - 1 Terminal (1T)")
        criar_tabela_config("2T", "Fast Charging - 2 Terminals (2T)")
        criar_tabela_config("Lenta", "Slow Charging (Overnight)")
        criar_tabela_config("Diesel", "Diesel Baseline (Combustion Engine)")

        for p in self.mista_1t_potencias_alvo:
            criar_tabela_config(f"M1T_{p}", f"Mixed Charging 1T - Fixed Power: {p} kW")
        for p in self.mista_2t_potencias_alvo:
            criar_tabela_config(f"M2T_{p}", f"Mixed Charging 2T - Fixed Power: {p} kW")

        doc.save(f"relatorio_graficos_completo_{L_usuario}m_{timestamp_agora}.docx")
        print("[SUCESSO WORD] Relatório com Novos Gráficos da Mista concluído!")


if __name__ == "__main__":
    app = GerenciadorRelatorio()
    app.executar()
